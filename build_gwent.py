"""Builds data/gwent.json from the Gwent sheet of the quest-order XLSX + tiva85 marker data."""
import sys, os, re, json, math, collections
from openpyxl import load_workbook
HERE=os.path.dirname(os.path.abspath(__file__)); D=os.path.join(HERE,'..','data')
XLSX=sys.argv[1] if len(sys.argv)>1 else os.path.join(HERE,'quest_order.xlsx')
ws=load_workbook(XLSX)['Gwent Cards']
MERGED={}
for rg in ws.merged_cells.ranges:
    v=ws.cell(rg.min_row,rg.min_col).value
    for r in range(rg.min_row,rg.max_row+1):
        for c in range(rg.min_col,rg.max_col+1): MERGED[(r,c)]=v
def cell(r,c):
    v=ws.cell(r,c).value
    return MERGED.get((r,c)) if v is None else v

# ---------- parse card copies ----------
cards=[]
for top,(r0,r1),heads in [('A',(3,47),['Northern Realms','Nilfgaard',"Scoia'tael"]),('B',(50,97),['Monsters','Skellige','Neutral'])]:
    for gi,(cc,lc) in enumerate([(2,5),(8,11),(14,17)]):
        name=None; prev=None
        for r in range(r0,r1+1):
            v=ws.cell(r,cc).value
            if isinstance(v,str) and v.strip() and v.strip()!='Card':
                if v.strip()!=name: prev=None
                name=re.sub(r'\s+',' ',v).strip()
            loc=cell(r,lc); chk=ws.cell(r,cc+1).value
            if (loc is None and chk is None) or name is None: continue
            src=re.sub(r'\s+',' ',str(loc or '')).strip() or prev or ''
            prev=src
            cards.append({'f':heads[gi],'c':name,'src':src})

# ---------- groups ----------
# vendor: (id, label, map, anchor signpost, snap categories, pick)  pick: index among nearest gwent-selling markers
VENDORS=[
 ('wo_inn','Elsa / Bram – hostinec, White Orchard','white_orchard',None,['innkeep'],0, r'Elsa or Bram'),
 ('wo_aldert','Aldert Geert, White Orchard','white_orchard',None,['gwent'],0, r'Aldert Geert'),
 ('crows_trader',"Obchodník – Crow's Perch",'hos_velen',"Crow's Perch",['shopkeeper'],0, r"Crow's Perch's trader"),
 ('crows_qm',"Proviantmajster – Crow's Perch",'hos_velen',"Crow's Perch",['gwent'],0, r"(?i)Crow's Perch's quartermaster"),
 ('claywich','Obchodník – Claywich','hos_velen','Claywich',['shopkeeper'],0, r"Claywich's merchant"),
 ('lindenvale','Obchodník – Lindenvale','hos_velen','Lindenvale',['shopkeeper'],0, r"Lindenvale's merchant"),
 ('midcopse','Obchodník – Midcopse','hos_velen','Midcopse',['shopkeeper'],0, r"Midcopse's merchant"),
 ('crossroads_inn','Krčmár – Inn at the Crossroads','hos_velen','Inn at the Crossroads',['innkeep'],0, r"Inn at the Crossroads'?s?'? ?(innkeep|or)|Lindenvale's innkeep"),
 ('stjepan','Stjepan – hostinec The Alchemy, Oxenfurt','hos_velen','Oxenfurt Harbor',['innkeep'],0, r'Stjepan at The Alchemy'),
 ('golden_sturgeon','Krčmár – The Golden Sturgeon, Novigrad','hos_velen','Hierarch Square',['innkeep'],'west', r'Golden Sturgeon'),
 ('kingfisher','Olivier – The Kingfisher, Novigrad','hos_velen','Hierarch Square',['innkeep'],'east', r'Kingfisher Inn'),
 ('cunny','Krčmár – Cunny of the Goose','hos_velen','Cunny of the Goose',['innkeep'],0, r'Cunny of the Goose'),
 ('seven_cats','Krčmár – Seven Cats Inn','hos_velen','Seven Cats Inn',['innkeep'],0, r'Seven Cats Inn'),
 ('passiflora','Marquise Serenity – Passiflora, Novigrad','hos_velen','Hierarch Square',['brothel'],0, r'Marquise Serenity at the Passiflora'),
 ('circus','Obchodník v cirkuse pri Carstene (HoS)','hos_velen','Carsten',['gwent','shopkeeper'],0, r'[Cc]ircus'),
 ('dulla',"Dulla kh'Amanni – Upper Mill (HoS)",'hos_velen','Upper Mill',['signpost'],0, r"Dulla kh'Amanni"),
 ('brunwich_barn','Stodola v Brunwichi (HoS)','hos_velen','Brunwich',['signpost'],0, r'barn in Brunwich'),
 ('vizima','Nilfgaardský šľachtic – palác vo Vizime',None,None,None,0, r'Vizima'),
 ('svorlag','Krčmár – Svorlag (Spikeroog)','skellige','Svorlag',['innkeep'],0, r"Svorlag's innkeep"),
 ('urialla','Krčmár – Urialla Harbor (An Skellig)','skellige','Urialla Harbor',['innkeep'],0, r"Urialla Harbor's [Ii]nnkeep"),
 ('harviken','Krčmár – Harviken (Faroe)','skellige','Harviken',['innkeep'],0, r"Harviken's innkeep"),
 ('arinbjorn','Krčmár – Arinbjorn','skellige','Arinbjorn',['innkeep'],0, r"Arinbjorn's innkeep|Innkeep in Arinbjorn"),
 ('jonas','Jonas – New Port Inn, Kaer Trolde','skellige','Kaer Trolde Harbor',['innkeep'],0, r'Jonas at [Tt]he New Port'),
 ('baw_ravello','Bylinkár – Castel Ravello','toussaint','Castel Ravello Vineyard',['herbalist','shopkeeper'],0, r'Castel Ravello'),
 ('baw_francollarts_arm','Zbrojár – Francollarts','toussaint','Francollarts',['armourer','shopkeeper'],0, r"Francollarts's armorer"),
 ('baw_cardinal','Krčmár – Scarlet Cardinal, Francollarts','toussaint','Francollarts',['innkeep'],0, r'Scarlet Cardinal'),
 ('baw_cockatrice','Krčmár – The Cockatrice Inn','toussaint','The Cockatrice Inn',['innkeep'],0, r'Cockatrice Inn'),
 ('baw_tourney','Tourney Grounds – krčmár, kováč, zbrojár, holič','toussaint','Tourney Grounds',['signpost'],0, r"Tourney Grounds?'s|Tourney Grounds' (armorer|barber)"),
 ('baw_isabelle','Madame Isabelle – prístav Beauclair','toussaint','Beauclair Port',['brothel'],0, r'Madame Isabelle'),
 ('baw_butcher','Mäsiar – prístav Beauclair','toussaint','Beauclair Port',['signpost'],0, r'Beauclair Port.s butcher'),
 ('baw_camerlengo','Vojvodský komorník – palác Beauclair','toussaint','Beauclair Palace',['signpost'],0, r'Camerlengo'),
 ('baw_beauclair','Beauclair – zbrojár, parfuméria, bylinkár, krajčír Pierre','toussaint',"Gran'Place",['signpost'],0, r"Beauclair's armorer|Perfumery|Herb Store|Pierre at Tailor"),
 ('baw_other','Toussaint – ďalší predajcovia (Barrel and Bung, Pheasantry, Dupont & Sons, Adder and Jewels)','toussaint',None,None,0, r'Barrel and Bung|Pheasantry|Dupont|Adder and Jewels'),
]
QUESTS=[
 ('q_velen','Gwent: Velen Players',r'Velen Players'),
 ('q_bigcity','Gwent: Big City Players',r'Big City Players'),
 ('q_innkeeps','Gwent: Playing Innkeeps',r'Playing Innkeeps'),
 ('q_oldpals','Gwent: Old Pals',r'Old Pals'),
 ('q_skellige','Gwent: Skellige Style',r'Skellige Style'),
 ('q_thaler','Gwent: Playing Thaler',r'Playing Thaler'),
 ('q_highstakes','High Stakes',r'High Stakes'),
 ('q_dangerous','A Dangerous Game',r'Dangerous Game'),
 ('q_lifedeath','A Matter of Life and Death',r'Matter of Life and Death'),
 ('q_shock','Shock Therapy',r'Shock Therapy'),
 ('q_thread','Following the Thread',r'Following the Thread'),
 ('q_dijkstra','Dijkstra (kúpeľný dom, Novigrad)',r'Sigismund Dijkstra'),
 ('q_hos','Hearts of Stone – Olgierd, Shani, Hilbert',r'Olgierd|Shani|Hilbert'),
 ('q_baw','Blood and Wine – turnaj a Louisova urna',r'Turn, Turn|Louis'),
]
M={mp:json.load(open(os.path.join(D,f'markers-{mp}.json'))) for mp in ['white_orchard','hos_velen','skellige','toussaint']}
def place(mp,anchor,snap,pick):
    data=M[mp]
    if snap==['signpost'] and anchor:
        s=[m for m in data['signpost'] if m[2].replace('*','')==anchor][0]; return s[0],s[1],False,None
    cand=[m for c in snap for m in data.get(c,[])]
    gw=[m for m in cand if 'Gwent' in m[3]] or cand
    if anchor is None: m=gw[0]; return m[0],m[1],False,m
    s=[m for m in data['signpost'] if m[2].replace('*','')==anchor][0]
    gw=sorted(gw,key=lambda m:math.dist((m[0],m[1]),(s[0],s[1])))
    if pick in ('west','east'):
        two=sorted(gw[:2],key=lambda m:m[1]); m=two[0] if pick=='west' else two[1]
    else: m=gw[pick]
    return m[0],m[1],math.dist((m[0],m[1]),(s[0],s[1]))>12,m
groups={}
for vid,label,mp,anchor,snap,pick,rx in VENDORS:
    g={'t':'vendor','n':label,'map':mp}
    if mp and snap:
        lat,lng,far,mk=place(mp,anchor,snap,pick); g.update(lat=round(lat,3),lng=round(lng,3))
        if mk is not None:
            txt=re.sub(r'<[^>]+>','',mk[3])
            w=re.findall(r"(Warning[^.]*\.|Not accessible after '[^']+' mission\.|May not be accessible[^.]*\.)",txt)
            if w: g['warn']=' '.join(w)
        if far or snap==['signpost']: g['approx']=True
    groups[vid]=g
QLINK={'q_velen':'Gwent: Velen Players','q_bigcity':'Gwent: Big City Players','q_innkeeps':'Gwent: Playing Innkeeps','q_oldpals':'Gwent: Old Pals',
 'q_skellige':'Gwent: Skellige Style','q_thaler':'Gwent: Playing Thaler','q_highstakes':'Gwent: High Stakes','q_dangerous':'A Dangerous Game',
 'q_lifedeath':'A Matter of Life and Death','q_shock':'Shock Therapy','q_thread':'Following the Thread'}
for qid,qn,rx in QUESTS:
    groups[qid]={'t':'quest','n':qn}
    if qid in QLINK: groups[qid]['quest']=QLINK[qid]
groups['vizima']['warn']='Iba počas prológu vo Vizime (Imperial Audience) – potom sa nedá získať.'
groups['random']={'t':'random','n':'Náhodné výhry od bežných hráčov'}
groups['other']={'t':'other','n':'Ostatné'}
out=[]; unmatched=collections.Counter(); base=0
for i,c in enumerate(cards):
    s=c['src']
    if s.lower().startswith('base deck'): base+=1; continue
    gs=[]
    if s=='Random': gs=['random']
    for vid,label,mp,anchor,snap,pick,rx in VENDORS:
        if re.search(rx,s): gs.append(vid)
    for qid,qn,rx in QUESTS:
        if re.search(rx,s):
            # "X's innkeep during Gwent: Playing Innkeeps" -> quest, not vendor
            gs=[g for g in gs if not (('during' in s or 'Won from' in s) and groups[g]['t']=='vendor')]+[qid]
    if not gs: gs=['other']; unmatched[s]+=1
    out.append({'id':'g%d'%len(out),'f':c['f'],'c':c['c'],'src':s,'g':sorted(set(gs),key=gs.index)})
json.dump({'cards':out,'groups':groups,'base':base},open(os.path.join(D,'gwent.json'),'w'),ensure_ascii=False,separators=(',',':'))
cnt=collections.Counter(g for c in out for g in c['g'])
print('cards',len(out),'base',base)
for k,v in groups.items(): print(f"{cnt[k]:3} {k:22} {v['n'][:45]:45} {v.get('map')} {'~' if v.get('approx') else ''} {v.get('warn','')} {v.get('quest','')}")
print('OTHER',dict(unmatched))
