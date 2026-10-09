import sys, os
from openpyxl import load_workbook
HERE=os.path.dirname(os.path.abspath(__file__)); OUT=os.path.join(HERE,'out')
XLSX=sys.argv[1] if len(sys.argv)>1 else os.path.join(HERE,'quest_order.xlsx')
import re, json, collections, unicodedata, difflib
wb=load_workbook(XLSX)
ws=wb['Order For MainSide Quests']
TYPE={'FFE06666':'main','FFF6B26B':'side','FFFFD966':'contract','FF93C47D':'treasure','FFA4C2F4':'pursuit','FF3D85C6':'scavenger','FF8E7CC3':'encounter'}
REGMAP={'KAER MORHEN':'kaer_morhen','WHITE ORCHARD':'white_orchard','VIZIMA':None,'VELEN':'hos_velen','NOVIGRAD/OXENFURT':'hos_velen','SKELLIGE':'skellige','HEARTS OF STONE':'hos_velen','BLOOD AND WINE':'toussaint'}
HL=re.compile(r'^=HYPERLINK\("([^"]+)",\s*"(.*)"\)\s*$',re.S)
def cellval(c):
    v=c.value; url=c.hyperlink.target if c.hyperlink else None
    if isinstance(v,str):
        m=HL.match(v.strip())
        if m: url=m.group(1); v=m.group(2)
        v=re.sub(r'[ \t]{2,}',' ',v).strip()
    return v,url
def fill(c): return c.fill.fgColor.rgb if c.fill.fill_type else None
items=[]; anytime=False; cur=None; region=None
for r in range(12, ws.max_row+1):
    A,B,D=ws.cell(r,1),ws.cell(r,2),ws.cell(r,4)
    a,_=cellval(A); b,burl=cellval(B); d,durl=cellval(D)
    if fill(A)=='FF434343':
        t=(a or '').strip()
        if t.upper().startswith('THE FOLLOWING QUESTS CAN BE DONE AT ANY TIME'):
            anytime=True; items.append({'kind':'anytime'})
        else:
            anytime=False
            if t: items.append({'kind':'note','text':t[0]+t[1:].lower() if t.isupper() else t})
        cur=None; continue
    if isinstance(b,str) and b.strip():
        m=re.match(r'^(.*?)\s*\((\d+)\)\s*$', b.strip())
        name=m.group(1).strip() if m else b.strip()
        lvl=int(m.group(2)) if m else None
        if isinstance(a,str) and a.strip(): region=a.strip()
        cur={'kind':'quest','row':r,'region':region,'name':name,'level':lvl,'type':TYPE.get(fill(B),'other'),'anytime':anytime,'wiki':burl,'notes':[]}
        items.append(cur)
        if d: cur['notes'].append({'t':str(d),'u':durl} if durl else {'t':str(d)})
    elif d and cur:
        cur['notes'].append({'t':str(d),'u':durl} if durl else {'t':str(d)})
    elif isinstance(a,str) and a.strip() and not b and fill(A)=='FF000000':
        items.append({'kind':'branch','text':a.strip().title()}); cur=None
# ---- marker matching
M=json.load(open(os.path.join(OUT,'markers_all.json')))
SCH={'cat':'fel','feline':'fel','griffin':'gri','wolf':'wol','wolven':'wol','ursine':'urs','bear':'urs','forgotten wolf':'fwol','manticore':'man'}
ROM={'i':'1','ii':'2','iii':'3','iv':'4','v':'5','vi':'6'}
def clean_label(s):
    s=re.sub(r'^\{.*?\}:\s*','',s)
    s=re.sub(r'<[^>]+>','',s)
    return s
def norm(s):
    s=clean_label(s).replace('´',"'").replace('’',"'").replace('`',"'"); s=unicodedata.normalize('NFKD',s).encode('ascii','ignore').decode().lower()
    s=re.sub(r'\(underground\)|\*','',s)
    s=re.sub(r'^(contract|scavenger hunt|gwent|races?|treasure hunt|vintner.?s contract|vitner.?s contract|the heroes. pursuits|fists of fury)\s*:\s*','',s)
    s=re.sub(r"[’']",'',s)
    s=re.sub(r'\(?part (\d+)\)?',r' p\1 ',s)
    s=re.sub(r'\(([ivx]+)\)',lambda m:' p'+ROM.get(m.group(1),m.group(1))+' ',s)
    s=re.sub(r"[^a-z0-9 ]",' ',s)
    s=re.sub(r'\b(the|a|an|of|s|crossings)\b',lambda m:'crossing' if m.group(1)=='crossings' else ' ',s)
    return re.sub(r'\s+',' ',s).strip()
def label_keys(lab):
    lab=clean_label(lab)
    keys={norm(lab)}
    head=re.split(r'\s+-\s+',lab)[0]
    m=re.match(r'^(.*?)\s+([IVX]+)$',head.strip())
    if m and m.group(2).lower() in ROM: head=f"{m.group(1)} (Part {ROM[m.group(2).lower()]})"
    keys.add(norm(head))
    return {k for k in keys if k}
def base(k): return re.sub(r'\s*\bp\d\b','',k).strip()
def scav_key(s):
    s=clean_label(s).lower()
    m=re.search(r'diagram\s*:\s*(feline|griffin|manticore|ursine|wolven)\s*-\s*grandmaster',s) or re.search(r'grandmaster (feline|griffin|manticore|ursine|wolven) gear',s)
    if m: return f"gm {SCH[m.group(1)]}"
    m=re.search(r'(forgotten wolf|cat|feline|griffin|wolf|wolven|ursine|bear) school( gear)?( upgrade)?( diagrams)?\s*-?\s*part (\d)',s)
    if m: return f"scav {SCH[m.group(1)]} {m.group(5)}"
    m=re.search(r'(forgotten wolf|cat|griffin|wolf|ursine|bear) school gear( upgrade)?( diagrams)?\s*(\(basic\))?\s*$',s)
    if m: return f"scav {SCH[m.group(1)]} basic"
    return None
idx=collections.defaultdict(list); bidx=collections.defaultdict(list)
for mp,dd in M.items():
    for cat,its in dd['cats'].items():
        for i,it in enumerate(its):
            it['label']=clean_label(it['label']); it['title']=clean_label(it['title'])
            keys=label_keys(it['label'])|label_keys(it['title'])
            sk=scav_key(it['label'])
            if sk: keys={sk}
            ref={'map':mp,'cat':cat,'i':i,'lat':it['lat'],'lng':it['lng']}
            for k in keys:
                idx[k].append(ref)
                if not sk: bidx[base(k)].append(ref)
json.dump(M,open(os.path.join(OUT,'markers_all.json'),'w'))
allkeys=list(idx)
stats=collections.Counter(); fuzzy=[]
for q in [x for x in items if x['kind']=='quest']:
    mp=REGMAP.get(q['region'])
    ALIAS={'Novigrad, Closed City I':'Novigrad, Closed City (Bandits)','Novigrad, Closed City II':'Novigrad, Closed City (Lussi, Fritz, Walter)'}
    sk=scav_key(q['name']); k=sk or norm(ALIAS.get(q['name'],q['name']))
    hits=idx.get(k,[])
    if not hits and not sk: hits=bidx.get(base(k),[])
    if not hits and q['type']!='main' and not sk:
        ptok=set(re.findall(r'\bp\d\b',k))
        cand=[c for c in difflib.get_close_matches(k,allkeys,n=3,cutoff=0.86) if set(re.findall(r'\bp\d\b',c))==ptok]
        if cand:
            hits=idx[cand[0]]; fuzzy.append((q['name'],cand[0]))
    # prefer markers on region map
    if mp and (q['type']=='main' or any(h['map']==mp for h in hits)): hits=[h for h in hits if h['map']==mp]
    # dedupe & cap
    seen=set(); hh=[]
    for h in hits:
        key=(h['map'],h['lat'],h['lng'])
        if key not in seen: seen.add(key); hh.append(h)
    dl=set(); act=set()
    for h in hh:
        pop=M[h['map']]['cats'][h['cat']][h['i']]['popup']
        dl.update(re.findall(r"Not accessible after '([^']+)' mission",pop)); act.update(re.findall(r"become active after the '([^']+)' mission",pop))
    if dl: q['until']=sorted(dl)
    if act: q['after']=sorted(act)
    q['markers']=hh[:12]; q['map']=mp if mp else (hh[0]['map'] if hh else None)
    stats[('hit' if hh else 'miss',q['type'])]+=1
for i,x in enumerate(items): x['id']=i
for x in items:
    if x['kind']=='quest':
        x['mk']=[[h['map'],h['cat'],h['i']] for h in x.pop('markers')]
json.dump(items,open(os.path.join(OUT,'quests.json'),'w'),ensure_ascii=False)
print('with until',sum(1 for x in items if x.get('until')), collections.Counter(u for x in items for u in x.get('until',[])))
print(sorted(stats.items()))
print('fuzzy',len(fuzzy)); [print('  ',a,'=>',b) for a,b in fuzzy]
miss=[x['name'] for x in items if x['kind']=='quest' and not x['mk'] and x['type']!='main']; print('MISS',miss)
print('total quests',sum(1 for x in items if x['kind']=='quest'), 'notes',sum(len(x.get('notes',[])) for x in items))
